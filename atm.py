
import os
from openpyxl import Workbook, load_workbook

FILE = "atm_accounts.xlsx"

# Dictionary to store account information
accounts = {}


# Load account data from Excel
def load_data():
    if os.path.exists(FILE):
        wb = load_workbook(FILE)
        ws = wb.active

        for row in ws.iter_rows(min_row=2, values_only=True):
            name, mobile, balance = row

            accounts[str(mobile)] = {
                "name": name,
                "mobile": str(mobile),
                "balance": float(balance)
            }

        wb.close()


# Save dictionary data to Excel
def save_data():
    wb = Workbook()
    ws = wb.active
    ws.title = "Accounts"

    ws.append(["Name", "Mobile Number", "Balance"])

    for account in accounts.values():
        ws.append([
            account["name"],
            account["mobile"],
            account["balance"]
        ])

    wb.save(FILE)
    wb.close()


# Create a new account
def create_account():
    name = input("Enter your name: ").strip()
    mobile = input("Enter your mobile number: ").strip()

    if not name or not mobile.isdigit() or len(mobile) != 10:
        print("Enter a valid name and 10-digit mobile number.")
        return

    if mobile in accounts:
        print("Account already exists!")
        return

    try:
        balance = float(input("Enter initial balance: "))

        if balance < 0:
            print("Balance cannot be negative.")
            return

    except ValueError:
        print("Enter a valid amount.")
        return

    accounts[mobile] = {
        "name": name,
        "mobile": mobile,
        "balance": balance
    }

    save_data()
    print("Your account is successfully created!")


# Add money
def add_money():
    mobile = input("Enter your mobile number: ").strip()

    if mobile not in accounts:
        print("Account not found!")
        return

    try:
        amount = float(input("Enter money: "))

        if amount <= 0:
            print("Enter a positive amount.")
            return

    except ValueError:
        print("Enter a valid amount.")
        return

    accounts[mobile]["balance"] += amount
    save_data()

    print("Your money successfully saved!")
    print("Updated balance: ₹", accounts[mobile]["balance"])


# Withdraw money
def withdraw_money():
    mobile = input("Enter your mobile number: ").strip()

    if mobile not in accounts:
        print("Account not found!")
        return

    try:
        amount = float(input("Enter amount to withdraw: "))

        if amount <= 0:
            print("Enter a positive amount.")
            return

    except ValueError:
        print("Enter a valid amount.")
        return

    balance = accounts[mobile]["balance"]

    if amount > balance:
        print("Insufficient balance!")
        print("Available balance: ₹", balance)
        return

    accounts[mobile]["balance"] -= amount
    save_data()

    print("Money successfully withdrawn!")
    print("Remaining balance: ₹", accounts[mobile]["balance"])


# Check account balance
def check_balance():
    mobile = input("Enter your mobile number: ").strip()

    if mobile in accounts:
        print("\n----- ACCOUNT INFORMATION -----")
        print("Name:", accounts[mobile]["name"])
        print("Mobile:", accounts[mobile]["mobile"])
        print("Current balance: ₹", accounts[mobile]["balance"])
    else:
        print("Account not found!")


# Main menu
load_data()

while True:
    print("\n===== COMMAND LINE ATM SYSTEM =====")
    print("1. Create Account")
    print("2. Add Money")
    print("3. Withdraw Money")
    print("4. Check Balance")
    print("5. Exit")

    choice = input("Enter your choice: ").strip()

    if choice == "1":
        create_account()

    elif choice == "2":
        add_money()

    elif choice == "3":
        withdraw_money()

    elif choice == "4":
        check_balance()

    elif choice == "5":
        print("Thank you for using our ATM system!")
        break

    else:
        print("Invalid choice! Please try again.")
